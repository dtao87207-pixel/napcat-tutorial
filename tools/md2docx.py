import os
import re, sys, io
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(HERE, "大肥鱼QQ机器人-NapCat方案教程.md")
OUT = os.path.join(HERE, "大肥鱼QQ机器人-NapCat方案教程.docx")

CN_FONT = "微软雅黑"
MONO_FONT = "Consolas"


def set_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(paragraph, hex_fill="F3F4F6"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    ppr.append(shd)


def add_inline(paragraph, text, base_size=10.5):
    # split on **bold**, `code`
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            inner = tok[2:-2]
            for piece in re.split(r"(`[^`]+`)", inner):
                if not piece:
                    continue
                if piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
                    r = paragraph.add_run(piece[1:-1])
                    set_font(r, MONO_FONT, base_size - 0.5, bold=True,
                             color=RGBColor(0xB9, 0x1C, 0x1C))
                else:
                    r = paragraph.add_run(piece)
                    set_font(r, CN_FONT, base_size, bold=True)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            set_font(r, MONO_FONT, base_size - 0.5, color=RGBColor(0xB9, 0x1C, 0x1C))
        else:
            r = paragraph.add_run(tok)
            set_font(r, CN_FONT, base_size)


def add_code_block(doc, lines):
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Inches(0.15)
        pf.line_spacing = 1.0
        shade(p)
        r = p.add_run(line if line else " ")
        set_font(r, MONO_FONT, 9)
    # spacer after block
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncol = len(header)
    table = doc.add_table(rows=1, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, cell_text in enumerate(header):
        cell = hdr[i]
        cell.text = ""
        p = cell.paragraphs[0]
        add_inline(p, re.sub(r"\*\*", "", cell_text), 9)
        for r in p.runs:
            r.font.bold = True
    for row in body:
        cells = table.add_row().cells
        for i in range(ncol):
            txt = row[i] if i < len(row) else ""
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_inline(p, re.sub(r"\*\*", "", txt), 9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def main():
    md = io.open(SRC, encoding="utf-8").read().splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = CN_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    i = 0
    in_code = False
    code_buf = []
    while i < len(md):
        line = md[i]

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, code_buf)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(md) and re.match(r"^\|[\s\-:|]+\|?$", md[i + 1].strip()):
            rows = [parse_table_row(stripped)]
            i += 2
            while i < len(md) and md[i].strip().startswith("|"):
                rows.append(parse_table_row(md[i].strip()))
                i += 1
            add_table(doc, rows)
            continue

        if stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run("─" * 40)
            set_font(r, CN_FONT, 8, color=RGBColor(0xD1, 0xD5, 0xDB))
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(14 if level <= 2 else 10)
            pf.space_after = Pt(6)
            if level == 1:
                r = p.add_run(text)
                set_font(r, CN_FONT, 20, bold=True, color=RGBColor(0x11, 0x18, 0x27))
            elif level == 2:
                r = p.add_run(text)
                set_font(r, CN_FONT, 15, bold=True, color=RGBColor(0x1F, 0x3B, 0x73))
            elif level == 3:
                r = p.add_run(text)
                set_font(r, CN_FONT, 12.5, bold=True, color=RGBColor(0x33, 0x41, 0x55))
            else:
                r = p.add_run(text)
                set_font(r, CN_FONT, 11, bold=True, color=RGBColor(0x37, 0x41, 0x51))
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(0.2)
            pf.space_after = Pt(4)
            add_inline(p, stripped.lstrip("> ").strip(), 10)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x55, 0x65, 0x7A)
            i += 1
            continue

        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_inline(p, stripped)
        i += 1

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
